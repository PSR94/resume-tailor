import base64
import io
import sys
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from main import app


client = TestClient(app)
ORIGINAL_BULLET = "Built reliable Python services for document processing."
REPLACEMENT_BULLET = "Improved document workflows by adding targeted Python automation."


def make_resume_base64() -> str:
    doc = Document()
    doc.add_heading("EXPERIENCE", level=1)

    role = doc.add_paragraph()
    run = role.add_run("Software Engineer")
    run.bold = True

    doc.add_paragraph(ORIGINAL_BULLET, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode()


def docx_text_from_base64(docx_base64: str) -> str:
    docx_bytes = base64.b64decode(docx_base64)
    doc = Document(io.BytesIO(docx_bytes))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def test_health_returns_ok():
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_extract_text_rejects_invalid_base64():
    response = client.post("/extract-text", json={"resume_base64": "not valid base64"})

    assert response.status_code == 400
    assert response.json()["detail"] == "Invalid base64 resume data."


def test_extract_text_rejects_invalid_docx_bytes():
    invalid_docx = base64.b64encode(b"not a docx").decode()

    response = client.post("/extract-text", json={"resume_base64": invalid_docx})

    assert response.status_code == 400
    assert response.json()["detail"] == "Uploaded file is not a valid DOCX."


def test_extract_text_accepts_minimal_valid_docx():
    response = client.post("/extract-text", json={"resume_base64": make_resume_base64()})

    assert response.status_code == 200
    data = response.json()
    assert "EXPERIENCE" in data["text"]
    assert ORIGINAL_BULLET in data["text"]


def test_apply_swaps_with_no_swaps_returns_docx_and_preserves_bullet_count():
    response = client.post(
        "/apply-swaps",
        json={
            "resume_base64": make_resume_base64(),
            "manifest": {
                "jobTitle": "Software Engineer",
                "company": "Local Test",
                "swaps": [],
            },
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["docx_base64"]
    assert data["swaps_applied"] == 0
    assert data["bullet_count_before"] == data["bullet_count_after"]
    assert data["bullet_count_before"] == 1


def test_apply_swaps_with_one_approved_swap_replaces_target_bullet():
    response = client.post(
        "/apply-swaps",
        json={
            "resume_base64": make_resume_base64(),
            "manifest": {
                "jobTitle": "Software Engineer",
                "company": "Local Test",
                "swaps": [
                    {
                        "id": "swap-1",
                        "action": "swap",
                        "section": "EXPERIENCE",
                        "roleIndex": 0,
                        "bulletIndex": 0,
                        "originalBullet": ORIGINAL_BULLET,
                        "newBullet": REPLACEMENT_BULLET,
                        "approved": True,
                    }
                ],
            },
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["docx_base64"]
    assert data["swaps_applied"] == 1
    assert data["swaps_requested"] == 1
    assert data["swaps_skipped"] == 0
    assert data["bullet_count_before"] == data["bullet_count_after"]

    returned_text = docx_text_from_base64(data["docx_base64"])
    assert REPLACEMENT_BULLET in returned_text
    assert ORIGINAL_BULLET not in returned_text


def test_apply_swaps_rejects_unmatched_approved_swap():
    response = client.post(
        "/apply-swaps",
        json={
            "resume_base64": make_resume_base64(),
            "manifest": {
                "jobTitle": "Software Engineer",
                "company": "Local Test",
                "swaps": [
                    {
                        "id": "swap-missing",
                        "action": "swap",
                        "section": "EXPERIENCE",
                        "roleIndex": 0,
                        "bulletIndex": 99,
                        "originalBullet": "This bullet is not in the resume.",
                        "newBullet": REPLACEMENT_BULLET,
                        "approved": True,
                    }
                ],
            },
        },
    )

    assert response.status_code == 422
    detail = response.json()["detail"]
    assert "1 approved swap(s) could not be applied" in detail
    assert "swap-missing: target bullet not found" in detail
